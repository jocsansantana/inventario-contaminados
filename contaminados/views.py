import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Sum
from .models import BultoContaminado
from .forms import BultoContaminadoForm

# --- VISTA PRINCIPAL (Formulario + Tabla) ---
def contaminados_view(request):
    if request.method == 'POST':
        form = BultoContaminadoForm(request.POST)
        if form.is_valid():
            bulto = form.save(commit=False)
            if request.user.is_authenticated:
                bulto.usuario_registro = request.user
            else:
                from django.contrib.auth.models import User
                bulto.usuario_registro = User.objects.first()
            bulto.save()
            messages.success(request, f"¡Bulto {bulto.codigo} registrado exitosamente!")
            return redirect('contaminados:tabla_bultos')
        else:
            messages.error(request, "Error en el formulario. Por favor verifica los datos.")
    else:
        form = BultoContaminadoForm()

    bultos = BultoContaminado.objects.all()
    total_bultos = bultos.count()
    total_peso_neto = bultos.aggregate(Sum('peso_neto_kg'))['peso_neto_kg__sum'] or 0

    context = {
        'form': form,
        'bultos': bultos,
        'total_bultos': total_bultos,
        'total_peso_neto': total_peso_neto,
    }
    return render(request,'tabla_bultos.html',context)

#Eliminar Bulto
def eliminar_bulto(request, pk):
    bulto = get_object_or_404(BultoContaminado, pk=pk)
    
    if request.method == 'POST':
        codigo_eliminar = bulto.codigo
        bulto.delete()  # ✅ Elimina físicamente el registro en MySQL
        messages.success(request, f'¡El bulto {codigo_eliminar} ha sido eliminado correctamente!')
    
    return redirect('contaminados:tabla_bultos')

# --- VISTA DE EXPORTACIÓN A EXCEL PRO (.XLSX) ---
def exportar_contaminados_excel(request):
    # 1. Consulta limpia sin select_related
    bultos = BultoContaminado.objects.all()

    # 2. Crear libro de trabajo openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bultos Contaminados"

    # Estilos del Excel
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='800020', end_color='800020', fill_type='solid') # Color Guinda
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # Encabezados de las columnas
    headers = [
        "Código / Lote", 
        "Tipo Material", 
        "Subtipo / Grado", 
        "Causa Contaminación", 
        "Fuente de Origen", 
        "Peso Neto (kg)", 
        "Observaciones", 
        "Fecha Registro"
    ]
    ws.append(headers)

    # Formato a la fila de encabezados
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    # 3. Llenado de datos
    for bulto in bultos:
        ws.append([
            bulto.codigo,
            bulto.get_tipo_material_display(),
            bulto.get_subtipo_material_display(),
            bulto.get_causa_contaminacion_display(),
            bulto.fuente_origen or 'N/A',  # Texto directo de fuente_origen
            float(bulto.peso_neto_kg) if bulto.peso_neto_kg else 0.0,
            bulto.observaciones or '',
            bulto.fecha_registro.strftime('%d/%m/%Y %H:%M') if bulto.fecha_registro else ''
        ])

    # Formato de celdas y bordes
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        row[0].alignment = align_left    # Código
        row[1].alignment = align_center  # Tipo
        row[2].alignment = align_center  # Subtipo
        row[3].alignment = align_center  # Causa
        row[4].alignment = align_center  # Fuente de Origen
        row[5].alignment = align_right   # Peso
        row[5].number_format = '#,##0.00'
        row[6].alignment = align_left    # Observaciones
        row[7].alignment = align_center  # Fecha

        for cell in row:
            cell.border = thin_border

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 4. Respuesta HTTP para descargar el archivo .xlsx
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Inventario_Contaminados.xlsx"'
    wb.save(response)

    return response